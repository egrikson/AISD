"""Преобразование отчёта из Markdown в документ Word (.docx).

Поддерживается подмножество Markdown, которое используется в отчёте:
заголовки `#`…`###`, абзацы, маркированные и нумерованные списки, таблицы с
разделителем `|`, блоки кода в тройных обратных кавычках, изображения
`![подпись](путь)`, разрыв страницы `\\pagebreak`, выделение `**жирным**` и
`` `моноширинным` ``.

Оформление приближено к требованиям к отчётам: Times New Roman 14 пт,
полуторный интервал, выравнивание по ширине, абзацный отступ 1.25 см.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
BODY_SIZE = Pt(14)
TABLE_SIZE = Pt(9)
CODE_SIZE = Pt(8)
MAX_IMAGE_WIDTH = Cm(16)

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


# --------------------------------------------------------------------------
#                          низкоуровневые помощники
# --------------------------------------------------------------------------
def _setup_styles(document: Document) -> None:
    """Базовый стиль документа: шрифт, размер, интервалы, поля."""
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    paragraph_format = normal.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.space_after = Pt(0)

    for section in document.sections:
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)


def _add_inline(paragraph, text: str, font: str = BODY_FONT, size=BODY_SIZE) -> None:
    """Добавляет текст с разбором `**жирного**`, `*курсива*` и `моноширинного`."""
    for piece in INLINE_PATTERN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(size.pt - 2)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        if run.font.name is None:
            run.font.name = font
        if run.font.size is None:
            run.font.size = size


def _add_paragraph(document: Document, text: str, first_line_indent: bool = True):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    _add_inline(paragraph, text)
    return paragraph


def _add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt({1: 18, 2: 16, 3: 14}.get(level, 14))


def _add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = CODE_FONT
        run.font.size = CODE_SIZE
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        if index < len(lines) - 1:
            run.add_break()


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, value, size=TABLE_SIZE)
            for run in paragraph.runs:
                run.font.size = TABLE_SIZE
                if i == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def _add_image(document: Document, path: Path, caption: str) -> None:
    if not path.exists():
        _add_paragraph(document, f"[изображение не найдено: {path.name}]")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=MAX_IMAGE_WIDTH)
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(caption)
        run.italic = True
        run.font.size = Pt(11)
        run.font.name = BODY_FONT


def _add_title_page(document: Document, student: dict[str, str]) -> None:
    """Титульный лист по стандартной форме."""
    def centered(text: str, size: int = 14, bold: bool = False, space_after: int = 0):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(space_after)
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        return paragraph

    centered("Министерство науки и высшего образования Российской Федерации", 12)
    centered(student["university"], 12, bold=True, space_after=18)
    centered(student["faculty"], 12)
    centered(student["department"], 12, space_after=100)

    centered("ОТЧЁТ", 20, bold=True, space_after=6)
    centered(f"по дисциплине «{student['subject']}»", 14, space_after=18)
    centered(student["work"], 16, bold=True, space_after=6)
    centered(f"«{student['topic']}»", 14, space_after=110)

    for text in (
        f"Выполнил: студент группы {student['group']}",
        student["student"],
        "",
        f"Проверил: {student['teacher']}",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.font.name = BODY_FONT
        run.font.size = Pt(14)

    for _ in range(3):
        document.add_paragraph()
    centered(student["city_year"], 14)
    document.add_page_break()


def _continuation_length(lines: list[str], index: int) -> int:
    """Сколько следующих строк являются продолжением текущего пункта списка."""
    count = 0
    while index + count + 1 < len(lines):
        nxt = lines[index + count + 1]
        if not nxt.startswith(("  ", "\t")) or not nxt.strip():
            break
        if re.match(r"^\s*([*-]|\d+\.)\s", nxt):
            break
        count += 1
    return count


def _join_continuation(lines: list[str], text: str, index: int) -> str:
    """Склеивает пункт списка с его строками-продолжениями."""
    for offset in range(1, _continuation_length(lines, index) + 1):
        text += " " + lines[index + offset].strip()
    return text


# --------------------------------------------------------------------------
#                              основной разбор
# --------------------------------------------------------------------------
def write_docx(markdown: str, output_path: Path, base_dir: Path,
               student: dict[str, str]) -> Path:
    """Преобразует Markdown-текст отчёта в файл .docx."""
    document = Document()
    _setup_styles(document)
    _add_title_page(document, student)

    lines = markdown.splitlines()
    index = 0
    skip_title_block = markdown.lstrip().startswith("<!--TITLEPAGE-->")

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # титульный блок из Markdown в docx не переносим -- он собран отдельно
        if skip_title_block:
            if stripped == "\\pagebreak":
                skip_title_block = False
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped == "\\pagebreak":
            document.add_page_break()
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index].rstrip())
                index += 1
            index += 1
            _add_code_block(document, block or [""])
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and set(
            lines[index + 1].strip().replace("|", "").replace(" ", "")
        ) <= {"-", ":"} and lines[index + 1].strip().startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if set("".join(cells)) <= {"-", ":", " "} and rows:
                    index += 1
                    continue
                rows.append(cells)
                index += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            _add_table(document, rows)
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", stripped)
        if image_match:
            caption = ""
            # следующая непустая строка вида *подпись* используется как подпись
            look = index + 1
            while look < len(lines) and not lines[look].strip():
                look += 1
            if look < len(lines):
                caption_match = re.fullmatch(r"\*(.+)\*", lines[look].strip())
                if caption_match:
                    caption = caption_match.group(1)
                    index = look
            _add_image(document, (base_dir / image_match.group(2)).resolve(), caption)
            index += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            _add_heading(document, heading_match.group(2), len(heading_match.group(1)))
            index += 1
            continue

        list_match = re.match(r"^([*-])\s+(.*)$", stripped)
        if list_match:
            text = list_match.group(2)
            text = _join_continuation(lines, text, index)
            index += _continuation_length(lines, index)
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, text)
            index += 1
            continue

        numbered_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered_match:
            text = numbered_match.group(2)
            text = _join_continuation(lines, text, index)
            index += _continuation_length(lines, index)
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, text)
            index += 1
            continue

        if stripped.startswith("<!--"):
            index += 1
            continue

        # обычный абзац: склеиваем последовательные строки
        paragraph_lines = [stripped]
        while index + 1 < len(lines):
            nxt = lines[index + 1].strip()
            if (not nxt or nxt.startswith(("#", "|", "```", "!", "* ", "- ", "\\pagebreak"))
                    or re.match(r"^\d+\.\s", nxt)):
                break
            paragraph_lines.append(nxt)
            index += 1
        _add_paragraph(document, " ".join(paragraph_lines))
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
